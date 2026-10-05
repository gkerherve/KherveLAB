from __future__ import annotations

import shutil
import subprocess
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest

from khervelab.core.config import booking_yaml, load_config
from khervelab.core.facility import FacilityService
from khervelab.core.models import Booking, User
from khervelab.core.reports import bookable_hours, build_report
from khervelab.reports_export import export_khervetex, export_xlsx, report_tex, tex

LON = ZoneInfo("Europe/London")
KHERVETEX = Path(__file__).resolve().parents[2] / "KherveTeX"
REAL_HOME = str(Path.home())  # captured before the fixtures point HOME at a temp dir


def seed(svc, inst, start, hours, user, kind="measurement"):
    b = Booking(inst, start, start + timedelta(hours=hours), user, kind=kind)
    rel = b.path(svc.tz).as_posix()
    svc.repo.write_file(rel, booking_yaml(b))
    svc.repo.commit("seed", [rel])


@pytest.fixture
def year(facility):
    """A year of history: past bookings are seeded directly (the rules
    refuse bookings in the past, as they should)."""
    svc = FacilityService(facility)
    svc.save_user(User("u-0002", "Surfaces group", permissions=("xps",)))
    svc.save_user(User("u-0003", "Catalysis group", permissions=("xps", "nap-xps")))
    for m in range(1, 13):
        seed(svc, "xps", datetime(2025, m, 10, 9, tzinfo=LON), 4, "u-0002")
        seed(svc, "xps", datetime(2025, m, 20, 9, tzinfo=LON), 6, "u-0003")
        seed(svc, "nap-xps", datetime(2025, m, 15, 9, tzinfo=LON), 8, "u-0003")
    seed(svc, "nap-xps", datetime(2025, 6, 2, 9, tzinfo=LON), 3, "u-0002", kind="training")
    seed(svc, "xps", datetime(2024, 12, 31, 22, tzinfo=LON), 4, "u-0002")  # straddles the start
    svc.reload()
    svc.log("xps", "fault", "ion gun failed", blocking=True,
            when=datetime(2025, 5, 1, 9, tzinfo=LON))
    (f,) = [x for x in __import__("khervelab.core.logbook", fromlist=["faults"]).faults(svc.cfg.logs)]
    svc.log("xps", "fault_update", "new gun fitted", fault=f.id, status="resolved",
            when=datetime(2025, 5, 3, 9, tzinfo=LON))
    return svc


def test_bookable_hours_handles_dst(facility):
    cfg = load_config(facility.path)
    allweek = cfg.instruments["xps"]
    # 2025 has 8760 hours; 30 March is 23 h, 26 October is 25 h
    assert bookable_hours(allweek, date(2025, 1, 1), date(2026, 1, 1), LON) == 8760
    assert bookable_hours(allweek, date(2025, 3, 30), date(2025, 3, 31), LON) == 23
    assert bookable_hours(allweek, date(2025, 10, 26), date(2025, 10, 27), LON) == 25
    weekday = cfg.instruments["sem-sigma-300"]   # 08:00-20:00 weekdays
    assert bookable_hours(weekday, date(2025, 10, 6), date(2025, 10, 13), LON) == 60


def test_full_year_numbers(year):
    rep = build_report(year.cfg, date(2025, 1, 1), date(2026, 1, 1))
    xps = next(i for i in rep.instruments if i.id == "xps")
    nap = next(i for i in rep.instruments if i.id == "nap-xps")
    assert xps.booked == pytest.approx(12 * 10 + 2)        # +2 h of the straddling booking
    assert nap.booked == pytest.approx(96) and nap.training == pytest.approx(3)
    assert xps.utilisation == pytest.approx(122 / 8760)
    assert xps.downtime == pytest.approx(48) and xps.faults == 1
    assert rep.by_group == {"Catalysis group": pytest.approx(168),
                            "Surfaces group": pytest.approx(53)}
    assert sum(rep.monthly["2025-03"].values()) == pytest.approx(18)
    assert xps.users == 2


def test_local_names_and_groups_override(year):
    rep = build_report(year.cfg, date(2025, 1, 1), date(2026, 1, 1),
                       names={"u-0002": "Ada Lovelace"}, groups={"u-0002": "Prof X group"})
    assert "Ada Lovelace" in rep.by_user and "Prof X group" in rep.by_group


def test_open_fault_counts_until_now(facility):
    svc = FacilityService(facility)
    now = datetime.now(timezone.utc)
    svc.log("xps", "fault", "leak", blocking=True, when=(now - timedelta(hours=10)).astimezone(LON))
    rep = build_report(svc.cfg, date.today() - timedelta(days=5), date.today() + timedelta(days=60))
    xps = next(i for i in rep.instruments if i.id == "xps")
    assert 9.9 < xps.downtime < 10.2


def test_xlsx_has_summary_and_raw_rows(year, tmp_path):
    from openpyxl import load_workbook
    rep = build_report(year.cfg, date(2025, 1, 1), date(2026, 1, 1))
    out = export_xlsx(rep, tmp_path / "r.xlsx")
    wb = load_workbook(out)
    assert {"Summary", "By group", "By user", "Monthly", "Downtime", "Training",
            "Raw bookings"} <= set(wb.sheetnames)
    assert wb["Raw bookings"].max_row == 1 + len(rep.usage)
    assert wb["Summary"]._charts and wb["Monthly"]._charts


def test_tex_escaping():
    assert tex("50% & #1_a {x}") == r"50\% \& \#1\_a \{x\}"


def test_tex_imports_into_khervetex(year, tmp_path):
    if not (KHERVETEX / "khervedoc" / "importers.py").exists():
        pytest.skip("KherveTeX checkout not found next to KherveLAB")
    rep = build_report(year.cfg, date(2025, 1, 1), date(2026, 1, 1))
    src = report_tex(rep)
    code = ("import sys; sys.path.insert(0, sys.argv[1]); "
            "from khervedoc.importers import import_tex; "
            "doc = import_tex(open(sys.argv[2]).read()); "
            "src = repr(doc.children); "
            "print(len(doc.children), src.count('tikzpicture'))")
    path = tmp_path / "r.tex"
    path.write_text(src)
    res = subprocess.run([sys.executable, "-c", code, str(KHERVETEX), str(path)],
                         capture_output=True, text=True)
    if "No module named" in res.stderr:
        pytest.skip(f"KherveTeX dependencies missing here: {res.stderr.strip().splitlines()[-1]}")
    assert res.returncode == 0, res.stderr
    blocks, charts = map(int, res.stdout.split())
    assert blocks > 5 and charts >= 3  # the pgfplots charts survive the import


@pytest.mark.skipif(shutil.which("tectonic") is None, reason="tectonic not installed")
def test_full_year_report_compiles(year, tmp_path):
    rep = build_report(year.cfg, date(2025, 1, 1), date(2026, 1, 1))
    tex_path = export_khervetex(rep, tmp_path / "proj", "report")
    import os
    env = {**os.environ, "HOME": REAL_HOME}  # tectonic's package cache lives under HOME

    def compile_(*extra):
        return subprocess.run(["tectonic", "-X", "compile", *extra, tex_path.name],
                              cwd=tex_path.parent, capture_output=True, text=True, timeout=600,
                              env=env)
    res = compile_("--only-cached")  # fast when the cache is warm
    if res.returncode != 0:
        res = compile_()
    assert res.returncode == 0, res.stderr[-2000:]
    assert (tex_path.parent / "report.pdf").read_bytes()[:4] == b"%PDF"


def test_notebook_api(year, tmp_path, monkeypatch):
    from khervelab import notebook as lab
    from khervelab.local.dataindex import DataIndex
    home = tmp_path / "home-kl"
    lab.remember(year.repo.path, home)
    s = year.add_sample("Pt foil")
    data = tmp_path / "data"
    data.mkdir()
    shutil.copy(Path(__file__).parent / "data" / "Pt4f.vms", data / f"{s.id}.vms")
    idx = DataIndex(home / "index.db")
    idx.add_folder(data, "xps")
    idx.scan()
    idx.match(year)
    idx.close()
    rows = lab.find("Pt 4f", home=home, as_frame=False)
    assert rows and rows[0]["sample"] == s.id
    info = lab.sample(s.id, home=home)
    assert info["files"] and info["lineage"] == [s.id]
    region = lab.spectra(rows[0]["path"])["Pt 4f"]
    be = region.binding_energy()
    assert len(region.y) == 281 and 61 < min(be) < 63 and 89 < max(be) < 91
    rep = lab.report(date(2025, 1, 1), date(2026, 1, 1), home=home)
    assert rep.total_booked > 0
