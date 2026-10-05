"""Usage and cost reports, per user or for the whole lab, as CSV, Excel and PDF.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

Only approved bookings are charged; a booking belongs to the period its
start falls in. Each booking carries the hourly rate in force when it was
made, so changing a rate never rewrites past charges.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from . import db, logic


@dataclass
class Line:
    booking_id: int
    user_id: int
    user: str
    username: str
    email: str
    group: str
    category: str
    instrument: str
    start: str
    end: str
    hours: float
    rate: float
    cost: float
    purpose: str
    session: str = ""
    price: float | None = None     # fixed session price, else charged by the hour

    def basis(self, currency: str) -> str:
        if self.price is not None:
            return f"{currency}{self.price:,.2f}/session"
        return f"{currency}{self.rate:,.2f}/h"

@dataclass
class Total:
    key: str
    hours: float = 0.0
    cost: float = 0.0
    bookings: int = 0


@dataclass
class Report:
    lab: str
    currency: str
    start: date
    end: date                      # inclusive, as people say "1 to 31 October"
    lines: list[Line] = field(default_factory=list)

    def _totals(self, attr: str) -> list[Total]:
        out: dict[str, Total] = {}
        for ln in self.lines:
            t = out.setdefault(getattr(ln, attr), Total(getattr(ln, attr)))
            t.hours += ln.hours
            t.cost += ln.cost
            t.bookings += 1
        return sorted(out.values(), key=lambda t: (-t.cost, t.key))

    def by_user(self) -> list[Total]:
        return self._totals("user")

    def by_group(self) -> list[Total]:
        return self._totals("group")

    def by_instrument(self) -> list[Total]:
        return self._totals("instrument")

    @property
    def total_cost(self) -> float:
        return round(sum(ln.cost for ln in self.lines), 2)

    @property
    def total_hours(self) -> float:
        return sum(ln.hours for ln in self.lines)

    def period(self) -> str:
        return f"{self.start:%d %b %Y} – {self.end:%d %b %Y}"


def build(conn, start: date, end: date, user_id: int | None = None,
          instrument_id: int | None = None) -> Report:
    q = ("SELECT b.*, u.full_name, u.username, u.email, u.group_name, u.category, "
         "i.name AS instrument, s.name AS session FROM bookings b JOIN users u ON u.id=b.user_id "
         "JOIN instruments i ON i.id=b.instrument_id LEFT JOIN sessions s ON s.id=b.session_id "
         "WHERE b.status='approved' AND b.start >= ? AND b.start < ?")
    args: list = [start.isoformat(), (end + timedelta(days=1)).isoformat()]
    if user_id:
        q += " AND b.user_id=?"
        args.append(user_id)
    if instrument_id:
        q += " AND b.instrument_id=?"
        args.append(instrument_id)
    q += " ORDER BY u.full_name, b.start"
    rep = Report(db.setting(conn, "lab_name"), db.setting(conn, "currency"), start, end)
    for b in conn.execute(q, args).fetchall():
        h = logic.hours(conn, logic.parse(b["start"]), logic.parse(b["end"]))
        rep.lines.append(Line(b["id"], b["user_id"], b["full_name"], b["username"], b["email"],
                              b["group_name"], b["category"], b["instrument"], b["start"],
                              b["end"], round(h, 2), b["rate"], logic.cost(conn, b),
                              b["purpose"], b["session"] or "", b["price"]))
    return rep


HEADER = ["Booking", "User", "Username", "Group", "Category", "Instrument", "Session", "Start",
          "End", "Hours", "Rate per hour", "Price per session", "Cost", "Purpose"]


def _row(ln: Line) -> list:
    return [ln.booking_id, ln.user, ln.username, ln.group, ln.category, ln.instrument,
            ln.session, ln.start.replace("T", " "), ln.end.replace("T", " "), ln.hours,
            "" if ln.price is not None else ln.rate,
            ln.price if ln.price is not None else "", ln.cost, ln.purpose]


def to_csv(rep: Report) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADER)
    for ln in rep.lines:
        w.writerow(_row(ln))
    w.writerow([])
    w.writerow(["", "Total", "", "", "", "", "", "", "", round(rep.total_hours, 2), "", "",
                rep.total_cost, ""])
    return buf.getvalue()


def to_xlsx(rep: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{rep.lab} — instrument usage and costs"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([rep.period()])
    ws.append([])
    for title, totals in (("By user", rep.by_user()), ("By group", rep.by_group()),
                          ("By instrument", rep.by_instrument())):
        ws.append([title, "Bookings", "Hours", f"Cost ({rep.currency})"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for t in totals:
            ws.append([t.key or "—", t.bookings, round(t.hours, 2), round(t.cost, 2)])
        ws.append([])
    ws.append(["Total", len(rep.lines), round(rep.total_hours, 2), rep.total_cost])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    detail = wb.create_sheet("Bookings")
    detail.append(HEADER)
    for c in detail[1]:
        c.font = Font(bold=True)
    for ln in rep.lines:
        detail.append(_row(ln))
    for col, w in zip("ABCDEFGHIJKLMN", (9, 24, 14, 20, 16, 26, 22, 17, 17, 8, 12, 14, 10, 40)):
        detail.column_dimensions[col].width = w
    detail.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rep: Report) -> bytes:
    """One statement per user (a page each), then a lab summary when the
    report covers several users."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    cur = rep.currency
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"{rep.lab} usage statement")
    grid = TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f4")),
                       ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                       ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                       ("ALIGN", (-3, 1), (-1, -1), "RIGHT"),
                       ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                       ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")])
    story: list = []
    users: dict[int, list[Line]] = defaultdict(list)
    for ln in rep.lines:
        users[ln.user_id].append(ln)
    for n, lines in enumerate(users.values()):
        u = lines[0]
        if n:
            story.append(PageBreak())
        story += [Paragraph(f"{rep.lab}", st["Title"]),
                  Paragraph(f"Instrument usage statement · {rep.period()}", st["Heading3"]),
                  Paragraph(f"<b>{u.user}</b> ({u.username})"
                            + (f" · {u.group}" if u.group else "")
                            + (f" · {u.email}" if u.email else "")
                            + (f" · rate category: {u.category}" if u.category else ""),
                            st["Normal"]), Spacer(1, 5 * mm)]
        rows = [["Date", "Instrument", "Time", "Hours", "Rate", f"Cost ({cur})"]]
        for ln in lines:
            s, e = logic.parse(ln.start), logic.parse(ln.end)
            when = f"{s:%H:%M}–{e:%H:%M}" + ("" if e.date() == s.date()
                                             else f" (+{(e.date() - s.date()).days} d)")
            rows.append([f"{s:%d %b %Y}", ln.instrument + (f" · {ln.session}" if ln.session
                                                           else ""), when,
                         f"{ln.hours:.2f}", ln.basis(cur), f"{ln.cost:,.2f}"])
        rows.append(["Total", "", "", f"{sum(l.hours for l in lines):.2f}", "",
                     f"{sum(l.cost for l in lines):,.2f}"])
        t = Table(rows, colWidths=[22 * mm, 58 * mm, 30 * mm, 15 * mm, 31 * mm, 22 * mm],
                  repeatRows=1)
        t.setStyle(grid)
        story.append(t)
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(f"Generated {datetime.now():%d %B %Y %H:%M} by KherveLAB.",
                               st["Italic"]))
    if len(users) != 1:
        if users:
            story.append(PageBreak())
        story += [Paragraph(f"{rep.lab}", st["Title"]),
                  Paragraph(f"Usage summary · {rep.period()}", st["Heading3"])]
        rows = [["User", "Group", "Bookings", "Hours", f"Cost ({cur})"]]
        groups = {ln.user: ln.group for ln in rep.lines}
        for t in rep.by_user():
            rows.append([t.key, groups.get(t.key, ""), str(t.bookings), f"{t.hours:.2f}",
                         f"{t.cost:,.2f}"])
        rows.append(["Total", "", str(len(rep.lines)), f"{rep.total_hours:.2f}",
                     f"{rep.total_cost:,.2f}"])
        t = Table(rows, colWidths=[55 * mm, 45 * mm, 22 * mm, 22 * mm, 30 * mm], repeatRows=1)
        t.setStyle(grid)
        story.append(t)
    doc.build(story)
    return buf.getvalue()
