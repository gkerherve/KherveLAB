"""Finance and usage statistics for a period, ready to tabulate and chart.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

One Analytics object feeds every view: the app's report tabs, the web
reports page, the PDF lab report and the Excel workbook. Charges always go
through logic.cost, so every report agrees with the statements.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from . import db, logic

WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")


@dataclass
class Chart:
    """A chart any renderer can draw: kind is bar, hbar, pie, line, stacked
    or heatmap; series are (name, values) aligned with labels."""
    kind: str
    title: str
    labels: list[str]
    series: list[tuple[str, list[float]]]
    unit: str = ""
    rows: list[str] = field(default_factory=list)      # heatmap row labels

    def table(self) -> list[list]:
        if self.kind == "heatmap":
            return [[""] + self.labels] + [[r] + list(v) for r, (_, v) in
                                           zip(self.rows, self.series)]
        return [[""] + [n for n, _ in self.series]] + [
            [lab] + [vals[i] for _, vals in self.series] for i, lab in enumerate(self.labels)]


@dataclass
class Kpi:
    label: str
    value: str
    note: str = ""


@dataclass
class Analytics:
    lab: str
    currency: str
    start: date
    end: date                         # inclusive
    kpis: list[Kpi] = field(default_factory=list)
    charts: dict[str, Chart] = field(default_factory=dict)
    instruments: list[dict] = field(default_factory=list)   # per-instrument table
    groups: list[dict] = field(default_factory=list)        # per-group table

    def period(self) -> str:
        return f"{self.start:%d %b %Y} – {self.end:%d %b %Y}"

    SECTIONS = {
        "Finance": ["revenue_month", "revenue_instrument", "revenue_category",
                    "revenue_group", "revenue_user", "price_basis"],
        "Usage": ["hours_month", "heatmap", "period_split", "status", "lead_time"],
        "Instruments": ["utilisation"],
        "Downtime": ["downtime"],
    }


def _clip_hours(conn, s: datetime, e: datetime, p0: datetime, p1: datetime) -> float:
    a, b = max(s, p0), min(e, p1)
    return logic.hours(conn, a, b) if b > a else 0.0


def _month_keys(start: date, end: date) -> list[str]:
    out, d = [], date(start.year, start.month, 1)
    while d <= end:
        out.append(f"{d:%Y-%m}")
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return out


def _period_kind(conn, inst, s: datetime) -> str:
    """Daytime / Evening / Weekend for free time; the session name otherwise."""
    if inst["booking_mode"] == "sessions":
        occ = [o for o in logic.occurrences(conn, inst["id"], s, s + timedelta(minutes=1))
               if o.start <= s < o.end]
        return occ[0].name if occ else "Other"
    ps = [p for p in logic.periods(inst, s, s + timedelta(minutes=1)) if p.start <= s < p.end]
    return ps[0].label if ps else "Outside hours"


def build(conn, start: date, end: date, user_id: int | None = None,
          instrument_id: int | None = None) -> Analytics:
    cur = db.setting(conn, "currency")
    a = Analytics(db.setting(conn, "lab_name"), cur, start, end)
    p0 = datetime(start.year, start.month, start.day)
    p1 = datetime(end.year, end.month, end.day) + timedelta(days=1)
    now = logic.now_local(conn)
    inst_q = "SELECT * FROM instruments" + (" WHERE id=?" if instrument_id else "") + \
        " ORDER BY name"
    insts = conn.execute(inst_q, (instrument_id,) if instrument_id else ()).fetchall()
    by_id = {i["id"]: i for i in insts}
    ids = list(by_id) or [-1]
    marks = ",".join("?" * len(ids))
    rows = conn.execute(
        "SELECT b.*, u.full_name, u.group_name, u.category, u.created AS user_created "
        f"FROM bookings b JOIN users u ON u.id=b.user_id WHERE b.instrument_id IN ({marks}) "
        "AND b.start >= ? AND b.start < ?" + (" AND b.user_id=?" if user_id else "")
        + " ORDER BY b.start",
        (*ids, logic.fmt(p0), logic.fmt(p1), *((user_id,) if user_id else ()))).fetchall()
    approved = [r for r in rows if r["status"] == "approved"]

    revenue = sum(logic.cost(conn, r) for r in approved)
    hours = sum(logic.hours(conn, logic.parse(r["start"]), logic.parse(r["end"]))
                for r in approved)
    pending = [r for r in rows if r["status"] == "pending"]
    pending_value = sum(logic.cost(conn, r) for r in pending)

    # -- per instrument: bookable, booked, utilisation, revenue, downtime ----------------
    downtime_down, downtime_problem = {}, {}
    for inst in insts:
        bookable = sum(_clip_hours(conn, s0, s1, p0, p1)
                       for s0, s1, _ in logic.slots(conn, inst, p0, p1))
        mine = [r for r in approved if r["instrument_id"] == inst["id"]]
        booked = sum(_clip_hours(conn, logic.parse(r["start"]), logic.parse(r["end"]), p0, p1)
                     for r in mine)
        dd = dp = 0.0
        for i in logic.issues(conn, inst["id"], p0, p1):
            e = logic.parse(i["end"]) if i["end"] else min(now, p1)
            h = _clip_hours(conn, logic.parse(i["start"]), e, p0, p1)
            if i["kind"] == "down":
                dd += h
            else:
                dp += h
        downtime_down[inst["name"]], downtime_problem[inst["name"]] = dd, dp
        a.instruments.append({
            "name": inst["name"], "bookings": len(mine), "hours": round(booked, 2),
            "bookable": round(bookable, 2),
            "utilisation": round(100 * booked / bookable, 1) if bookable else 0.0,
            "revenue": round(sum(logic.cost(conn, r) for r in mine), 2),
            "down_hours": round(dd, 1), "problem_hours": round(dp, 1),
            "users": len({r["user_id"] for r in mine})})
    total_bookable = sum(i["bookable"] for i in a.instruments)

    # -- headline figures --------------------------------------------------------------------
    active = {r["user_id"] for r in approved}
    new_users = conn.execute("SELECT COUNT(*) FROM users WHERE created >= ? AND created < ?",
                             (logic.fmt(p0), logic.fmt(p1))).fetchone()[0]
    cancelled = sum(1 for r in rows if r["status"] in ("cancelled", "rejected"))
    a.kpis = [
        Kpi("Revenue", f"{cur}{revenue:,.2f}", f"{len(approved)} approved bookings"),
        Kpi("Hours booked", f"{hours:,.1f} h",
            f"{100 * hours / total_bookable:.1f}% of bookable time" if total_bookable else ""),
        Kpi("Active users", str(len(active)), f"{new_users} new account(s)"),
        Kpi("Waiting for approval", f"{cur}{pending_value:,.2f}", f"{len(pending)} request(s)"),
        Kpi("Cancelled or rejected", str(cancelled),
            f"{100 * cancelled / len(rows):.0f}% of requests" if rows else ""),
        Kpi("Out of order", f"{sum(downtime_down.values()):,.1f} h",
            f"problems reported: {sum(downtime_problem.values()):,.1f} h"),
    ]

    # -- finance -----------------------------------------------------------------------------
    months = _month_keys(start, end)
    rev_m, hrs_m = defaultdict(float), defaultdict(float)
    by_inst, by_cat, by_group, by_user, basis = (defaultdict(float) for _ in range(5))
    group_rows: dict[str, dict] = {}
    for r in approved:
        c = logic.cost(conn, r)
        s = logic.parse(r["start"])
        h = logic.hours(conn, s, logic.parse(r["end"]))
        rev_m[f"{s:%Y-%m}"] += c
        hrs_m[f"{s:%Y-%m}"] += h
        by_inst[by_id[r["instrument_id"]]["name"]] += c
        by_cat[r["category"] or "—"] += c
        g = r["group_name"] or "No group"
        by_group[g] += c
        by_user[r["full_name"]] += c
        basis["Fixed session price" if r["price"] is not None else "Hourly rate"] += c
        gr = group_rows.setdefault(g, {"group": g, "users": set(), "bookings": 0, "hours": 0.0,
                                       "revenue": 0.0})
        gr["users"].add(r["user_id"])
        gr["bookings"] += 1
        gr["hours"] += h
        gr["revenue"] += c
    a.groups = sorted(({**g, "users": len(g["users"]), "hours": round(g["hours"], 2),
                        "revenue": round(g["revenue"], 2)} for g in group_rows.values()),
                      key=lambda g: -g["revenue"])

    def ranked(d, limit=None):
        items = sorted(d.items(), key=lambda kv: -kv[1])
        if limit and len(items) > limit:
            rest = sum(v for _, v in items[limit:])
            items = items[:limit] + [("Others", rest)]
        return [k for k, _ in items], [round(v, 2) for _, v in items]

    mlabels = [datetime.strptime(m, "%Y-%m").strftime("%b %y") for m in months]
    a.charts["revenue_month"] = Chart("bar", "Revenue by month", mlabels,
                                      [("Revenue", [round(rev_m[m], 2) for m in months])], cur)
    lab, val = ranked(by_inst)
    a.charts["revenue_instrument"] = Chart("hbar", "Revenue by instrument", lab,
                                           [("Revenue", val)], cur)
    lab, val = ranked(by_cat)
    a.charts["revenue_category"] = Chart("pie", "Revenue by user category", lab,
                                         [("Revenue", val)], cur)
    lab, val = ranked(by_group, 8)
    a.charts["revenue_group"] = Chart("pie", "Revenue by group", lab, [("Revenue", val)], cur)
    lab, val = ranked(by_user, 15)
    a.charts["revenue_user"] = Chart("hbar", "Revenue by user (top 15)", lab,
                                     [("Revenue", val)], cur)
    lab, val = ranked(basis)
    a.charts["price_basis"] = Chart("pie", "Revenue by price basis", lab, [("Revenue", val)],
                                    cur)

    # -- usage ---------------------------------------------------------------------------------
    a.charts["hours_month"] = Chart("line", "Hours booked by month", mlabels,
                                    [("Hours", [round(hrs_m[m], 1) for m in months])], "h")
    heat = [[0.0] * 24 for _ in range(7)]
    split: dict[str, Counter] = defaultdict(Counter)
    for r in approved:
        s, e = logic.parse(r["start"]), logic.parse(r["end"])
        t = s.replace(minute=0)
        while t < e:
            t1 = t + timedelta(hours=1)
            heat[t.weekday()][t.hour] += _clip_hours(conn, max(t, s), min(t1, e), p0, p1)
            t = t1
        inst = by_id[r["instrument_id"]]
        split[inst["name"]][_period_kind(conn, inst, s)] += logic.hours(conn, s, e)
    a.charts["heatmap"] = Chart("heatmap", "When the lab is used (hours booked)",
                                [f"{h:02d}" for h in range(24)],
                                [(WEEKDAYS[d], [round(x, 1) for x in heat[d]]) for d in range(7)],
                                "h", rows=list(WEEKDAYS))
    kinds = sorted({k for c in split.values() for k in c},
                   key=lambda k: ("Daytime", "Evening", "Weekend").index(k)
                   if k in ("Daytime", "Evening", "Weekend") else 9)
    names = [i["name"] for i in insts if i["name"] in split]
    a.charts["period_split"] = Chart(
        "stacked", "Daytime, evening and weekend use (hours)", names,
        [(k, [round(split[n][k], 1) for n in names]) for k in kinds], "h")
    status = Counter(r["status"] for r in rows)
    order = ["approved", "pending", "rejected", "cancelled"]
    a.charts["status"] = Chart("pie", "What happened to booking requests",
                               [s.capitalize() for s in order if status[s]],
                               [("Bookings", [status[s] for s in order if status[s]])], "")
    lead: dict[str, list[float]] = defaultdict(list)
    for r in approved:
        if r["created"]:
            days = (logic.parse(r["start"]) - logic.parse(r["created"])).total_seconds() / 86400
            lead[by_id[r["instrument_id"]]["name"]].append(max(0.0, days))
    names = sorted(lead)
    a.charts["lead_time"] = Chart("hbar", "How far ahead people book (average days)", names,
                                  [("Days", [round(sum(lead[n]) / len(lead[n]), 1)
                                             for n in names])], "days")

    # -- instruments and downtime ----------------------------------------------------------------
    names = [i["name"] for i in a.instruments]
    a.charts["utilisation"] = Chart("hbar", "Utilisation (booked ÷ bookable hours)", names,
                                    [("Utilisation", [i["utilisation"] for i in a.instruments])],
                                    "%")
    a.charts["downtime"] = Chart("stacked", "Downtime and problems (hours)", names,
                                 [("Out of order", [round(downtime_down[n], 1) for n in names]),
                                  ("Problem", [round(downtime_problem[n], 1) for n in names])],
                                 "h")
    return a


# -- exports --------------------------------------------------------------------------------

def _tables(a: Analytics) -> list[tuple[str, list[list]]]:
    cur = a.currency
    inst = [["Instrument", "Bookings", "Users", "Hours booked", "Bookable hours",
             "Utilisation %", f"Revenue ({cur})", "Out of order h", "Problem h"]]
    inst += [[i["name"], i["bookings"], i["users"], i["hours"], i["bookable"], i["utilisation"],
              i["revenue"], i["down_hours"], i["problem_hours"]] for i in a.instruments]
    grp = [["Group", "Users", "Bookings", "Hours", f"Revenue ({cur})"]]
    grp += [[g["group"], g["users"], g["bookings"], g["hours"], g["revenue"]] for g in a.groups]
    return [("Instruments", inst), ("Groups", grp)]


def to_pdf(a: Analytics) -> bytes:
    """A lab report: headline figures, every chart, and the instrument and
    group tables."""
    import io

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (CondPageBreak, KeepTogether, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    from . import charts

    st = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=f"{a.lab} lab report")
    width = A4[0] - 32 * mm
    story: list = [Paragraph(a.lab, st["Title"]),
                   Paragraph(f"Lab report · {a.period()}", st["Heading3"]), Spacer(1, 4 * mm)]
    cells = [[Paragraph(f"<font size=8 color='#6b7785'>{k.label}</font><br/>"
                        f"<font size=15><b>{k.value}</b></font><br/>"
                        f"<font size=7.5 color='#6b7785'>{k.note}</font>",
                        ParagraphStyle("kpi", parent=st["Normal"], leading=17))
              for k in a.kpis[r:r + 3]] for r in range(0, len(a.kpis), 3)]
    kt = Table(cells, colWidths=[width / 3] * 3)
    kt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#dde3ea")),
                            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dde3ea")),
                            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f7f9fc")),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [kt, Spacer(1, 6 * mm)]
    grid = TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f4")),
                       ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                       ("FONTSIZE", (0, 0), (-1, -1), 8),
                       ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                       ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey)])
    for section, keys in Analytics.SECTIONS.items():
        story += [CondPageBreak(90 * mm), Paragraph(section, st["Heading2"])]
        for k in keys:
            story += [KeepTogether([charts.to_drawing(a.charts[k], width)]), Spacer(1, 4 * mm)]
        if section == "Instruments":
            title, rows = _tables(a)[0]
            t = Table([[str(c) for c in r] for r in rows], repeatRows=1)
            t.setStyle(grid)
            story.append(t)
    story += [PageBreak(), Paragraph("Groups", st["Heading2"])]
    t = Table([[str(c) for c in r] for r in _tables(a)[1][1]], repeatRows=1)
    t.setStyle(grid)
    story += [t, Spacer(1, 6 * mm),
              Paragraph(f"Generated {datetime.now():%d %B %Y %H:%M} by KherveLAB. Revenue counts "
                        "approved bookings at the price in force when each was made.",
                        st["Italic"])]
    doc.build(story)
    return buf.getvalue()


def to_xlsx(a: Analytics) -> bytes:
    """One sheet per chart with its numbers and a native Excel chart, plus
    the headline figures and the instrument and group tables."""
    import io

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws.append([f"{a.lab} — lab report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([a.period()])
    ws.append([])
    for k in a.kpis:
        ws.append([k.label, k.value, k.note])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 34
    for title, rows in _tables(a):
        t = wb.create_sheet(title)
        for r in rows:
            t.append(r)
        for c in t[1]:
            c.font = Font(bold=True)
        t.column_dimensions["A"].width = 24
        t.freeze_panes = "A2"
    for keys in Analytics.SECTIONS.values():
        for key in keys:
            ch = a.charts[key]
            s = wb.create_sheet(key.replace("_", " ").capitalize())
            for r in ch.table():
                s.append(r)
            for c in s[1]:
                c.font = Font(bold=True)
            s.column_dimensions["A"].width = 26
            n = len(ch.labels) if ch.kind != "heatmap" else len(ch.rows)
            if not n:
                continue
            if ch.kind == "heatmap":
                s.conditional_formatting.add(
                    f"B2:Y{n + 1}", ColorScaleRule(start_type="min", start_color="F7F9FC",
                                                   end_type="max", end_color="3B7DDD"))
                continue
            if ch.kind == "pie":
                xc = PieChart()
            elif ch.kind == "line":
                xc = LineChart()
            else:
                xc = BarChart()
                xc.type = "bar" if ch.kind in ("hbar", "stacked") else "col"
                if ch.kind == "stacked":
                    xc.grouping, xc.overlap = "stacked", 100
            xc.title = ch.title
            xc.add_data(Reference(s, min_col=2, max_col=1 + len(ch.series), min_row=1,
                                  max_row=n + 1), titles_from_data=True)
            xc.set_categories(Reference(s, min_col=1, min_row=2, max_row=n + 1))
            xc.width, xc.height = 18, 9 + (n * 0.4 if ch.kind in ("hbar", "stacked") else 0)
            s.add_chart(xc, f"{chr(ord('B') + len(ch.series) + 1)}2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
