"""Import a lab's history from PPMS exports (CSV or Excel).

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.

PPMS installations name their export columns differently, so each kind of
file has a list of fields with the header names they usually go by;
guess_mapping matches them and the manager can correct the match. run()
does the whole import in one transaction and rolls it back for a dry run,
so the preview counts are exactly what an import would do.

Imported bookings are history: they skip the booking rules, keep the amount
PPMS charged as a fixed price, and imported accounts have no password until
their owner claims them (logic.claim_account).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from . import db, logic

NOTE = "Imported from PPMS"


@dataclass(frozen=True)
class Field:
    key: str
    label: str
    names: tuple[str, ...]
    required: bool = False


KINDS: dict[str, tuple[str, list[Field]]] = {
    "systems": ("Systems (instruments)", [
        Field("name", "Name", ("system", "system name", "systemname", "instrument", "name",
                               "equipment", "resource"), True),
        Field("type", "Type", ("type", "system type", "category", "core", "facility")),
        Field("description", "Description", ("description", "comments", "comment", "details")),
    ]),
    "users": ("Users", [
        Field("login", "Login", ("login", "username", "user login", "user name", "userlogin",
                                 "user id", "account"), True),
        Field("first", "First name", ("first name", "firstname", "fname", "given name")),
        Field("last", "Last name", ("last name", "lastname", "lname", "surname",
                                    "family name")),
        Field("full_name", "Full name", ("full name", "fullname", "name", "user")),
        Field("email", "Email", ("email", "e-mail", "email address", "mail")),
        Field("group", "Group", ("group", "group name", "unit", "lab", "pi", "group head",
                                 "head", "department", "team")),
        Field("category", "User type", ("user type", "usertype", "group type", "type",
                                        "affiliation", "category", "price type", "customer type",
                                        "ext/int", "internal/external")),
        Field("active", "Active", ("active", "status", "enabled", "account status")),
    ]),
    "prices": ("Prices", [
        Field("system", "System", ("system", "system name", "instrument", "name"), True),
        Field("category", "User type", ("user type", "group type", "type", "category",
                                        "price type", "affiliation"), True),
        Field("rate", "Hourly price", ("price per hour", "hourly price", "hourly rate",
                                       "rate", "price", "cost", "amount"), True),
    ]),
    "rights": ("Rights (training)", [
        Field("user", "Login", ("login", "username", "user login", "user"), True),
        Field("system", "System", ("system", "system name", "instrument"), True),
        Field("right", "Right", ("right", "rights", "status", "level", "autonomy", "access")),
    ]),
    "bookings": ("Bookings / usage", [
        Field("system", "System", ("system", "system name", "instrument", "resource"), True),
        Field("user", "Login", ("login", "username", "user login", "user", "booked by"), True),
        Field("date", "Date (if separate)", ("date", "day", "booking date", "session date")),
        Field("start", "Start", ("start", "start time", "starttime", "from", "begin", "start date",
                                 "booked from", "session start", "start date time"), True),
        Field("end", "End", ("end", "end time", "endtime", "to", "until", "stop", "end date",
                             "booked to", "session end", "end date time")),
        Field("hours", "Duration (h)", ("duration", "hours", "booked hours", "length",
                                        "duration h", "time used", "used hours")),
        Field("amount", "Amount charged", ("amount", "charge", "charged", "total", "price",
                                           "cost", "invoice amount", "final amount", "fee")),
        Field("status", "Status", ("status", "state", "cancelled", "canceled")),
        Field("project", "Project / account", ("project", "account", "cost centre",
                                                "cost center", "cost code", "grant", "fund")),
        Field("comment", "Comment", ("comment", "comments", "purpose", "description", "note",
                                     "assistance")),
        Field("created", "Booked on", ("booked on", "created", "booking made", "date booked",
                                       "creation date")),
    ]),
    "incidents": ("Incidents", [
        Field("system", "System", ("system", "system name", "instrument"), True),
        Field("start", "Start", ("start", "date", "reported", "from", "start date", "opened"),
              True),
        Field("end", "End", ("end", "resolved", "closed", "to", "end date", "fixed")),
        Field("severity", "Severity", ("severity", "level", "priority", "type", "status")),
        Field("description", "Description", ("description", "comment", "comments", "details",
                                             "problem", "subject", "title")),
    ]),
}
ORDER = ["systems", "users", "prices", "rights", "bookings", "incidents"]


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


@dataclass
class Table:
    path: str
    headers: list[str]
    rows: list[dict]


def read_table(path: str | Path) -> Table:
    """A CSV (any common delimiter or encoding) or the first sheet of an
    Excel file. The header row is the first with two or more filled cells,
    so a title line above it is skipped."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        grid = list(csv.reader(io.StringIO(text), dialect))
    start = next((i for i, r in enumerate(grid)
                  if sum(1 for c in r if c not in (None, "")) >= 2), 0)
    if not grid:
        return Table(str(path), [], [])
    headers, seen = [], set()
    for n, h in enumerate(grid[start]):
        h = str(h).strip() if h not in (None, "") else f"Column {n + 1}"
        while h in seen:
            h += " (2)"
        seen.add(h)
        headers.append(h)
    rows = []
    for r in grid[start + 1:]:
        if not any(c not in (None, "") for c in r):
            continue
        rows.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers)})
    return Table(str(path), headers, rows)


def guess_kind(headers: list[str]) -> str | None:
    """Which kind of export a file looks like, from its headers."""
    hs = {_norm(h) for h in headers}
    score = {}
    for kind, (_, fields) in KINDS.items():
        mapping = guess_mapping(kind, headers)
        req = [f for f in fields if f.required]
        if all(mapping.get(f.key) for f in req):
            hits = len([v for v in mapping.values() if v])
            score[kind] = (hits, hits / len(fields))
    if "bookings" in score and not ({"start", "start time", "from", "start date"} & hs):
        score.pop("bookings")
    return max(score, key=score.get) if score else None


def guess_mapping(kind: str, headers: list[str]) -> dict[str, str | None]:
    """field key → the column it most likely is, each column used once:
    exact names first, then a header containing a name."""
    norm = {h: _norm(h) for h in headers}
    used: set[str] = set()
    out: dict[str, str | None] = {}
    fields = KINDS[kind][1]
    for exact in (True, False):
        for f in fields:
            if out.get(f.key):
                continue
            for name in f.names:
                hit = next((h for h in headers if h not in used and
                            (norm[h] == name if exact else
                             re.search(rf"\b{re.escape(name)}\b", norm[h]))), None)
                if hit:
                    out[f.key] = hit
                    used.add(hit)
                    break
    return {f.key: out.get(f.key) for f in fields}


# -- value parsing --------------------------------------------------------------------------

_DT_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",
               "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M",
               "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d %b %Y %H:%M", "%d-%b-%Y %H:%M",
               "%d %B %Y %H:%M", "%d/%m/%y %H:%M", "%Y/%m/%d %H:%M", "%Y/%m/%d %H:%M:%S"]
_D_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%d %b %Y", "%d-%b-%Y",
              "%d %B %Y", "%d/%m/%y", "%Y/%m/%d"]


def _month_first(fmts):
    return [f.replace("%d/%m", "%m/%d") for f in fmts]


def parse_datetime(v, dayfirst: bool = True) -> datetime | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.replace(second=0, microsecond=0, tzinfo=None)
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = re.sub(r"\s+", " ", str(v).strip())
    s = re.sub(r"(\d)\s*(am|pm)$", lambda m: m.group(1) + " " + m.group(2).upper(), s,
               flags=re.I)
    fmts = _DT_FORMATS if dayfirst else _month_first(_DT_FORMATS)
    dfmts = _D_FORMATS if dayfirst else _month_first(_D_FORMATS)
    for f in fmts + [x + " %p" for x in fmts if "%H:%M" in x] + dfmts:
        try:
            d = datetime.strptime(s, f.replace("%H", "%I") if f.endswith("%p") else f)
            return d.replace(second=0)
        except ValueError:
            continue
    return None


def parse_time(v) -> time | None:
    if isinstance(v, time):
        return v.replace(second=0, microsecond=0)
    if isinstance(v, datetime):
        return v.time().replace(second=0, microsecond=0)
    m = re.fullmatch(r"\s*(\d{1,2})[:h.](\d{2})(?::\d{2})?\s*(am|pm)?\s*", str(v or ""), re.I)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if m.group(3):
        h = h % 12 + (12 if m.group(3).lower() == "pm" else 0)
    return time(h % 24, mi) if h <= 24 and mi < 60 else None


def parse_number(v) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else \
            s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", "") if re.search(r",\d{3}$", s) else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_hours(v) -> float | None:
    if isinstance(v, timedelta):
        return v.total_seconds() / 3600
    if isinstance(v, time):
        return v.hour + v.minute / 60
    m = re.fullmatch(r"\s*(\d+):(\d{2})(?::\d{2})?\s*", str(v or ""))
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60
    return parse_number(v)


def _yes(v) -> bool | None:
    s = _norm(v)
    if not s:
        return None
    if s in ("0", "no", "n", "false", "f", "inactive", "disabled", "deactivated", "d",
             "locked", "closed", "expired"):
        return False
    return True


def match_category(value, cats: list[str]) -> str | None:
    """A lab category for a PPMS user or price type, or None when nothing
    fits (the caller then adds the PPMS name as a new category)."""
    s = _norm(value)
    if not s:
        return None
    for c in cats:
        if _norm(c) == s:
            return c
    rules = [(("industr", "commerc", "private", "company", "corporate", "external non"),
              ("industr", "commerc")),
             (("extern", "academ", "other univ", "outside"), ("extern",)),
             (("intern", "inhouse", "in house", "local"), ("intern",))]
    for needles, targets in rules:
        if any(n in s for n in needles):
            hit = next((c for c in cats if any(t in _norm(c) for t in targets)), None)
            if hit:
                return hit
    return None


def booking_status(v) -> str:
    s = _norm(v)
    if any(w in s for w in ("cancel", "delet", "removed")):
        return "cancelled"
    if "reject" in s or "refus" in s or "denied" in s:
        return "rejected"
    if any(w in s for w in ("pending", "await", "request", "tentative")):
        return "pending"
    return "approved"


def incident_kind(v) -> str:
    s = _norm(v)
    return "down" if any(w in s for w in ("down", "out of order", "critical", "blocking",
                                          "broken", "high", "severe", "major")) else "problem"


def right_trained(v) -> bool:
    s = _norm(v)
    if not s:
        return True
    if s in ("n", "novice", "d", "deactivated", "no", "none", "0", "false", "pending",
             "requested", "trainee", "training"):
        return False
    return True


# -- the import -----------------------------------------------------------------------------

@dataclass
class Source:
    kind: str
    table: Table
    mapping: dict[str, str | None]

    def get(self, row: dict, key: str):
        col = self.mapping.get(key)
        v = row.get(col) if col else None
        return v.strip() if isinstance(v, str) else v


@dataclass
class Summary:
    counts: dict[str, int] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)
    backup: str = ""

    def add(self, key: str, n: int = 1):
        self.counts[key] = self.counts.get(key, 0) + n

    def problem(self, msg: str):
        if len(self.problems) < 200:
            self.problems.append(msg)
        self.add("rows skipped")

    def text(self) -> str:
        order = ["instruments created", "instruments matched", "users created", "users matched",
                 "users created from bookings", "categories added", "prices set",
                 "trained users set", "bookings imported", "already imported",
                 "overlap existing bookings", "incidents imported", "rows skipped"]
        keys = [k for k in order if self.counts.get(k)] + \
            [k for k in self.counts if k not in order]
        lines = [f"{self.counts[k]:>7,}  {k}" for k in keys] or ["Nothing to import."]
        if self.problems:
            lines += ["", "Rows skipped:"] + [f"  • {p}" for p in self.problems]
        return "\n".join(lines)


_COLOURS = ["#1f6feb", "#8250df", "#cf222e", "#e16f24", "#9a6700", "#1a7f37", "#0e8a96",
            "#bf3989"]


def run(conn, sources: list[Source], dry_run: bool = True, dayfirst: bool = True) -> Summary:
    s = Summary()
    conn.execute("BEGIN IMMEDIATE")
    try:
        _import(conn, sorted(sources, key=lambda x: ORDER.index(x.kind)), s, dayfirst)
    except Exception:
        conn.rollback()
        raise
    if dry_run:
        conn.rollback()
    else:
        conn.commit()
    return s


def _import(conn, sources: list[Source], s: Summary, dayfirst: bool):
    insts = {_norm(r["name"]): r["id"] for r in conn.execute("SELECT id, name FROM instruments")}
    users = {}
    for r in conn.execute("SELECT id, username, email FROM users"):
        users[_norm(r["username"])] = r["id"]
    emails = {_norm(r["email"]): r["id"] for r in
              conn.execute("SELECT id, email FROM users WHERE email != ''")}
    cats = db.categories(conn)
    now = logic.fmt(logic.now_local(conn))

    def category(value) -> str:
        if value in (None, ""):
            return cats[0] if cats else ""
        hit = match_category(value, cats)
        if hit:
            return hit
        cats.append(str(value).strip())
        db.set_setting(conn, "categories", "\n".join(cats))
        s.add("categories added")
        return cats[-1]

    def instrument(name, description="", made_by=""):
        key = _norm(name)
        if key in insts:
            return insts[key]
        cur = conn.execute(
            "INSERT INTO instruments (name, description, colour, approval, open_time, "
            "close_time, weekends, max_minutes) VALUES (?, ?, ?, 'trained', '00:00', '24:00', "
            "1, 1440)", (str(name).strip(), description or "",
                         _COLOURS[len(insts) % len(_COLOURS)]))
        insts[key] = cur.lastrowid
        s.add("instruments created" + made_by)
        return cur.lastrowid

    def user(login, row=None, src=None):
        key = _norm(login)
        if key in users:
            return users[key]
        if src is not None:
            email = src.get(row, "email") or ""
            if _norm(email) in emails and email:
                users[key] = emails[_norm(email)]
                s.add("users matched")
                return users[key]
            first, last = src.get(row, "first") or "", src.get(row, "last") or ""
            full = src.get(row, "full_name") or f"{first} {last}".strip() or str(login)
            group = src.get(row, "group") or ""
            cat = category(src.get(row, "category"))
            status = "disabled" if _yes(src.get(row, "active")) is False else "active"
            made_by = ""
        else:
            email, full, group, cat, status = "", str(login), "", category(None), "active"
            made_by = " from bookings"
        cur = conn.execute(
            "INSERT INTO users (username, password_hash, full_name, email, group_name, category,"
            " role, status, created) VALUES (?, '', ?, ?, ?, ?, 'user', ?, ?)",
            (str(login).strip(), str(full).strip(), str(email).strip(), str(group).strip(), cat,
             status, now))
        users[key] = cur.lastrowid
        if email:
            emails[_norm(email)] = cur.lastrowid
        s.add("users created" + made_by)
        return cur.lastrowid

    for src in sources:
        for n, row in enumerate(src.table.rows, start=1):     # as numbered in the preview
            where = f"{Path(src.table.path).name} row {n}"
            if src.kind == "systems":
                name = src.get(row, "name")
                if not name:
                    s.problem(f"{where}: no system name")
                    continue
                if _norm(name) in insts:
                    s.add("instruments matched")
                else:
                    desc = " · ".join(str(x) for x in (src.get(row, "description"),
                                                       src.get(row, "type")) if x)
                    instrument(name, desc)
            elif src.kind == "users":
                login = src.get(row, "login")
                if not login:
                    s.problem(f"{where}: no login")
                    continue
                if _norm(login) in users:
                    s.add("users matched")
                else:
                    user(login, row, src)
            elif src.kind == "prices":
                name, rate = src.get(row, "system"), parse_number(src.get(row, "rate"))
                if not name or rate is None:
                    s.problem(f"{where}: needs a system and a price")
                    continue
                conn.execute("INSERT OR REPLACE INTO rates (instrument_id, category, rate) "
                             "VALUES (?, ?, ?)", (instrument(name), category(
                                 src.get(row, "category")), rate))
                s.add("prices set")
            elif src.kind == "rights":
                login, name = src.get(row, "user"), src.get(row, "system")
                if not login or not name:
                    s.problem(f"{where}: needs a login and a system")
                    continue
                if not right_trained(src.get(row, "right")):
                    continue
                conn.execute("INSERT OR IGNORE INTO authorised (user_id, instrument_id) "
                             "VALUES (?, ?)", (user(login), instrument(name)))
                s.add("trained users set")
            elif src.kind == "bookings":
                _booking(conn, src, row, where, s, dayfirst, instrument, user)
            elif src.kind == "incidents":
                name = src.get(row, "system")
                start = parse_datetime(src.get(row, "start"), dayfirst)
                if not name or start is None:
                    s.problem(f"{where}: needs a system and a start date")
                    continue
                end = parse_datetime(src.get(row, "end"), dayfirst)
                if end is not None and end <= start:
                    end = start + timedelta(days=1)
                conn.execute(
                    "INSERT INTO issues (instrument_id, kind, start, end, note, reported_by, "
                    "created) VALUES (?, ?, ?, ?, ?, NULL, ?)",
                    (instrument(name), incident_kind(src.get(row, "severity")),
                     logic.fmt(start), logic.fmt(end) if end else None,
                     f"{src.get(row, 'description') or ''} ({NOTE})".strip(), now))
                s.add("incidents imported")


def _booking(conn, src, row, where, s, dayfirst, instrument, user):
    name, login = src.get(row, "system"), src.get(row, "user")
    if not name or not login:
        s.problem(f"{where}: needs a system and a login")
        return
    day = parse_datetime(src.get(row, "date"), dayfirst) if src.mapping.get("date") else None

    def when(key):
        v = src.get(row, key)
        if day is not None:
            t = parse_time(v)
            if t is not None:
                return datetime.combine(day.date(), t)
        return parse_datetime(v, dayfirst)

    start = when("start")
    end = when("end") if src.mapping.get("end") else None
    hours = parse_hours(src.get(row, "hours")) if src.mapping.get("hours") else None
    if start is None:
        s.problem(f"{where}: cannot read the start {src.get(row, 'start')!r}")
        return
    if end is not None and end <= start and day is not None:
        end += timedelta(days=1)                     # a run past midnight
    if end is None and hours:
        end = start + timedelta(hours=hours)
    if end is None or end <= start:
        s.problem(f"{where}: no end time or duration")
        return
    amount = parse_number(src.get(row, "amount")) if src.mapping.get("amount") else None
    purpose = " · ".join(str(x) for x in (src.get(row, "comment"),
                                          f"Project: {src.get(row, 'project')}"
                                          if src.get(row, "project") else None) if x)
    created = parse_datetime(src.get(row, "created"), dayfirst) if src.mapping.get("created") \
        else None
    result = record_history(conn, instrument(name), user(login), start, end,
                            booking_status(src.get(row, "status")), amount, purpose, created)
    s.add({"exists": "already imported", "overlap": "overlap existing bookings"}.get(
        result, "bookings imported"))
    if result == "overlap":
        s.add("bookings imported")


def record_history(conn, iid: int, uid: int, start: datetime, end: datetime,
                   status: str = "approved", amount: float | None = None, purpose: str = "",
                   created: datetime | None = None, note: str = NOTE) -> str:
    """Store a booking that already happened elsewhere, as it was: no booking
    rules, and the amount charged (if known) as its fixed price. Returns
    'exists' (skipped), 'overlap' (stored, but it overlaps another booking)
    or 'ok'."""
    st, en = logic.fmt(start), logic.fmt(end)
    if end <= start:
        raise ValueError("the end must be after the start")
    if status not in ("approved", "pending", "rejected", "cancelled"):
        raise ValueError(f"unknown status {status!r}")
    if conn.execute("SELECT 1 FROM bookings WHERE instrument_id=? AND user_id=? AND start=? "
                    "AND end=?", (iid, uid, st, en)).fetchone():
        return "exists"
    overlap = status in ("approved", "pending") and conn.execute(
        "SELECT 1 FROM bookings WHERE instrument_id=? AND status IN ('approved','pending') "
        "AND start < ? AND end > ?", (iid, en, st)).fetchone()
    rate = conn.execute("SELECT r.rate FROM rates r JOIN users u ON u.category=r.category "
                        "WHERE r.instrument_id=? AND u.id=?", (iid, uid)).fetchone()
    conn.execute(
        "INSERT INTO bookings (instrument_id, user_id, start, end, purpose, status, rate, "
        "created, note, price) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (iid, uid, st, en, purpose, status, rate[0] if rate else 0.0,
         logic.fmt(created) if created else st, note, amount))
    return "overlap" if overlap else "ok"
